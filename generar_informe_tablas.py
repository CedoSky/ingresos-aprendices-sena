#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de Informe del Sistema de Control de Ingresos SENA - VERSIÓN CON TABLAS
Estructura modular con tablas grandes y bien definidas
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime

def add_heading_styled(doc, text, level):
    """Agrega un encabezado con estilo personalizado"""
    heading = doc.add_heading(text, level=level)
    heading.style = f'Heading {level}'
    
    if level == 1:
        for run in heading.runs:
            run.font.color.rgb = RGBColor(26, 82, 0)
            run.font.size = Pt(28)
            run.font.bold = True
    elif level == 2:
        for run in heading.runs:
            run.font.color.rgb = RGBColor(57, 169, 0)
            run.font.size = Pt(18)
    
    return heading

def add_paragraph_justified(doc, text, size=11, color=None):
    """Agrega párrafo justificado"""
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    for run in p.runs:
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    
    return p

def shade_cell(cell, color):
    """Agrega sombreado a una celda"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._element.get_or_add_tcPr().append(shading_elm)

def create_big_table(doc, data, header_color='1a5200', col_widths=None):
    """Crea tabla grande y bien formateada"""
    rows = len(data)
    cols = len(data[0]) if data else 1
    
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Light Grid Accent 1'
    
    # Encabezado
    header_cells = table.rows[0].cells
    for i, cell in enumerate(header_cells):
        shade_cell(cell, header_color)
        
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(11)
            for run in paragraph.runs:
                run.font.bold = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Populate header
    for i, header in enumerate(data[0]):
        header_cells[i].text = str(header)
    
    # Datos
    for row_idx in range(1, rows):
        row_cells = table.rows[row_idx].cells
        for col_idx in range(cols):
            cell = row_cells[col_idx]
            cell.text = str(data[row_idx][col_idx])
            
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    
    # Ancho de columnas
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
    
    return table

def main():
    doc = Document()
    
    # ==== PORTADA ====
    title = doc.add_heading('Sistema de Control de Ingresos SENA', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(26, 82, 0)
        run.font.size = Pt(36)
    
    subtitle = doc.add_paragraph('Informe Completo con Estructura en Tablas')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(57, 169, 0)
    
    doc.add_paragraph()
    
    info_text = doc.add_paragraph(
        f'Centro de Formación SENA\n'
        f'Versión: 6.0\n'
        f'Fecha: {datetime.now().strftime("%d de %B de %Y")}'
    )
    info_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in info_text.runs:
        run.font.size = Pt(12)
    
    doc.add_page_break()
    
    # ===============================================
    # MÓDULO 1: INGRESO CON CÓDIGO DE BARRAS
    # ===============================================
    
    add_heading_styled(doc, 'MÓDULO 1: SISTEMA DE INGRESO CON CÓDIGO DE BARRAS', 1)
    
    add_paragraph_justified(doc,
        'El módulo de ingreso es donde los aprendices registran su entrada y salida diaria escaneando su documento. '
        'Sistema simple, rápido y automático.')
    
    doc.add_paragraph()
    
    # Tabla: Flujo General
    add_heading_styled(doc, 'Flujo General del Sistema', 2)
    
    flujo_data = [
        ['Número', 'Paso', 'Descripción', 'Tiempo'],
        ['1', 'Acceder a portal', 'Persona se acerca a pantalla de ingreso con logo SENA', '5 seg'],
        ['2', 'Escanear documento', 'Usa lector de código de barras para capturar cédula', '2 seg'],
        ['3', 'Búsqueda en BD', 'Sistema verifica si documento existe en base de datos', '1 seg'],
        ['4', 'Marcar acceso', 'Sistema decide si es ENTRADA o SALIDA según historial', 'Automático'],
        ['5', 'Guardar con hora', 'Registra en BD con timestamp exacto (año/mes/día/hora/minuto)', 'Automático'],
        ['6', 'Mostrar confirmación', 'Pantalla muestra resultado (éxito/error) con foto y nombre', '2 seg'],
    ]
    
    create_big_table(doc, flujo_data, '1a5200', [1.0, 2.0, 4.5, 1.5])
    doc.add_paragraph()
    
    # Tabla: Pasos Detallados
    add_heading_styled(doc, 'Pasos Detallados para Ingresar', 2)
    
    pasos_data = [
        ['PASO', 'ACCIÓN', 'DETALLES', 'RESULTADO'],
        [
            '1',
            'Ubicarse en la pantalla',
            'Dirígete a la entrada donde verás una computadora/tablet con pantalla del sistema. Verás: logo SENA, cuadro de texto, instrucciones "Escanea tu documento".',
            'Pantalla visible y lista'
        ],
        [
            '2',
            'Preparar documento',
            'Ten tu cédula o carnet a mano. Limpia el código de barras si está sucio. El código está en la parte trasera como líneas verticales negras.',
            'Documento listo'
        ],
        [
            '3',
            'Escanear código',
            'Usa el lector de código de barras (dispositivo rojo/negro). Acerca tu cédula al escáner. No necesitas presionar nada. El lector capturará automáticamente. Escucharás un beep.',
            'Sistema procesando'
        ],
        [
            '4',
            'Esperar búsqueda',
            'El sistema tardará 1-2 segundos en consultar la base de datos. Verás "Procesando..." o un ícono de carga en la pantalla.',
            'Búsqueda en progreso'
        ],
        [
            '5',
            'Ver resultado',
            'La pantalla mostrará:\n- ✅ INGRESO EXITOSO (si es primera vez del día)\n- ✅ SALIDA EXITOSA (si ya habías ingresado)\n- Tu nombre completo, foto, programa y hora exacta.',
            'Acceso registrado'
        ],
        [
            '6',
            'Pasar adelante',
            'Ya tu entrada/salida está guardada en el sistema. Puedes proceder normalmente. No necesitas hacer nada más.',
            '✅ Listo'
        ],
    ]
    
    create_big_table(doc, pasos_data, '1a5200', [0.8, 1.5, 3.5, 1.7])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Mensajes del Sistema
    add_heading_styled(doc, 'Mensajes del Sistema: Significados y Acciones', 2)
    
    mensajes_data = [
        ['MENSAJE', 'SIGNIFICADO', 'QUÉ SIGNIFICA', 'ACCIÓN A TOMAR'],
        [
            '✅ INGRESO\nEXITOSO',
            'Entrada registrada\nprimer acceso del día',
            'Tu documento fue encontrado en el sistema y es tu primera vez en el día. Tu entrada está registrada con hora exacta.',
            'Puedes proceder normalmente. Ya estás dentro.'
        ],
        [
            '✅ SALIDA\nEXITOSA',
            'Salida registrada',
            'Tu documento fue encontrado y ya habías ingresado hoy. Ahora tu salida está registrada completando tu registro del día.',
            'Completaste tu registro. Puedes retirarte.'
        ],
        [
            '⚠️ NO\nENCONTRADO',
            'Documento no en BD',
            'Tu documento no existe en la base de datos. Tu información aún no fue registrada en el sistema.',
            'Ve a vigilancia para registro manual. Necesitarán tus datos personales.'
        ],
        [
            '⚠️ ACCESO\nDENEGADO',
            'Cuenta desactivada',
            'Tu cuenta fue desactivada por: vacaciones, suspensión, baja del programa, sanción o cambio de estado.',
            'Contacta con administración. Hay una restricción en tu acceso.'
        ],
        [
            '⚠️ ERROR DE\nLECTURA',
            'Código ilegible',
            'El código de barras de tu documento no se pudo leer. Posiblemente está dañado, sucio, rayado o roto.',
            'Limpia tu documento e intenta de nuevo. Si persiste, ve a vigilancia para registro manual.'
        ],
        [
            '⚠️ ERROR DE\nSISTEMA',
            'Problema técnico',
            'Hay un problema de conexión a internet o fallo en el servidor. El sistema no puede consultar la BD.',
            'Intenta de nuevo en 10 segundos. Si sigue fallando, reporta a vigilancia.'
        ],
    ]
    
    create_big_table(doc, mensajes_data, '1a5200', [1.2, 1.5, 2.5, 2.3])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Requisitos
    add_heading_styled(doc, 'Requisitos para que el Sistema Funcione', 2)
    
    requisitos_data = [
        ['REQUISITO', 'DESCRIPCIÓN', 'POR QUÉ ES IMPORTANTE', 'QUIÉN LO VERIFICA'],
        [
            '📋 Estar Registrado',
            'Tu información debe estar cargada en la BD: nombre, documento, programa, especialidad, foto.',
            'Sin registro previo, el sistema no te encontrará al escanear tu código.',
            'Administración'
        ],
        [
            '📞 Código Legible',
            'Tu cédula que código de barras debe estar en buen estado: sin rayones, manchas, roturas o deterioro.',
            'Si está dañado, el lector no puede capturarlo y necesitarás registro manual.',
            'Tú mismo (mantenlo limpio)'
        ],
        [
            '🌐 Internet Activo',
            'La computadora de ingreso debe tener conexión a internet (WiFi o ethernet) activa y estable.',
            'Sin conexión, el sistema no puede consultar la BD ni registrar acceso.',
            'Vigilancia/Administración'
        ],
        [
            '⚡ Hardware OK',
            'Computadora, monitor, lector de código y escritorio deben estar funcionando correctamente.',
            'Si algo falla (pantalla no enciende, lector no funciona), no se puede registrar.',
            'Vigilancia/Técnico'
        ],
        [
            '🕐 Hora Real',
            'La computadora debe tener la hora correcta del sistema (sincronizada automáticamente).',
            'Si la hora es incorrecta, los registros aparecerán con tiempo equivocado.',
            'Sistema automático'
        ],
    ]
    
    create_big_table(doc, requisitos_data, '1a5200', [1.3, 2.0, 2.5, 1.7])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Problemas y Soluciones
    add_heading_styled(doc, 'Problemas Comunes y Soluciones', 2)
    
    problemas_data = [
        ['PROBLEMA', 'POSIBLE CAUSA', 'CÓMO SOLUCIONARLO', 'SI NO FUNCIONA'],
        [
            'El lector no escanea',
            'Código sucio, rayado, dañado o lector desconectado',
            '1. Limpia tu documento con paño suave\n2. Intenta nuevamente\n3. Verifica que lector emita beep',
            'Ve a vigilancia para registro manual'
        ],
        [
            'Sale "Documento no encontrado"',
            'Tu información no fue registrada aún en el sistema',
            '1. Acude a administración\n2. Lleva tu cédula y datos personales\n3. Te registrarán en el sistema',
            'Pregunta por estado de tu solicitud'
        ],
        [
            'Sale "Acceso Denegado" sin razón',
            'Tu cuenta fue desactivada (error, cambio de estado, sanción)',
            '1. Contacta a administración\n2. Solicita revisión de tu estado\n3. Pide que reactiven tu acceso',
            'Espera confirmación de administración'
        ],
        [
            'Pantalla no responde o se cuelga',
            'Problema técnico, desconexión de internet, sobrecarga',
            '1. Espera 10 segundos\n2. Intenta de nuevo\n3. Si sigue igual, reporta a vigilancia',
            'Vigilancia reiniciará el sistema'
        ],
        [
            'Lector no funciona (sin beep)',
            'Dispositivo desconectado, batería agotada, falla de hardware',
            '1. Reporta a vigilancia\n2. Ellos verificarán el lector\n3. registro manual mientras se repara',
            'Espera reparación del dispositivo'
        ],
    ]
    
    create_big_table(doc, problemas_data, '1a5200', [1.5, 1.8, 2.5, 1.7])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # ===============================================
    # MÓDULO 2: VIGILANCIA
    # ===============================================
    
    add_heading_styled(doc, 'MÓDULO 2: PANEL DE VIGILANCIA', 1)
    
    add_paragraph_justified(doc,
        'El Panel de Vigilancia es la pantalla de control en tiempo real. El vigilante monitorea quién está adentro, '
        'detecta anomalías y controla accesos.')
    
    doc.add_paragraph()
    
    # Tabla: Cómo Acceder
    add_heading_styled(doc, 'Cómo Acceder al Panel de Vigilancia', 2)
    
    acceso_vig_data = [
        ['PASO', 'ACCIÓN', 'DETALLES'],
        [
            '1',
            'Abrir navegador',
            'Chrome, Firefox, Edge o Safari. Ve a la dirección del sistema (ej: http://localhost)'
        ],
        [
            '2',
            'Login como Vigilante',
            'Ingresa Email: vigilante01@sena.edu.co (tu usuario)\nContraseña: [tu contraseña segura]\nClick INGRESAR'
        ],
        [
            '3',
            'Ver Dashboard',
            'Se abre menú principal con opciones. Busca botón "CONTROL DE ACCESOS" o "VIGILANCIA"'
        ],
        [
            '4',
            'Panel se Carga',
            'El Panel de Vigilancia se abre. Espera 2-5 segundos a que cargue toda la información en tiempo real.'
        ],
        [
            '5',
            '✅ Listo',
            'Panel activo. Ves todos los registros, estadísticas, alertas. Sistema se actualiza automáticamente.'
        ],
    ]
    
    create_big_table(doc, acceso_vig_data, '1a5200', [0.8, 2.0, 3.7])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Componentes del Panel
    add_heading_styled(doc, 'Componentes del Panel de Vigilancia', 2)
    
    componentes_data = [
        ['COMPONENTE', 'UBICACIÓN', 'QUÉ MUESTRA', 'PARA QUÉ SIRVE'],
        [
            '📌 TOPBAR\n(Barra Superior)',
            'Parte superior\nde la pantalla',
            '• Logo SENA\n• Nombre panel: "Control de Accesos"\n• RELOJ DIGITAL con hora exacta\n• Nombre de vigilante conectado\n• Botón SALIR',
            'Verificar hora sincronizada\nIdentificar quién está conectado\nCerrar sesión de forma segura'
        ],
        [
            '📊 ESTADÍSTICAS\n(Cards Superiores)',
            'Debajo del topbar\nen 4 tarjetas',
            '• Ingresos Hoy: número total\n• Personas Dentro Ahora: cantidad actual\n• Faltas por Salir: quién debe irse\n• Alertas Activas: problemas detectados',
            'Ver resumen rápido del día\nIdentificar situaciones anormales\nEstar atento a alertas'
        ],
        [
            '🏢 ZONAS/PISOS\n(Panel Izquierdo)',
            'Columna izquierda\nde la pantalla',
            'Distribución de personas:\n• Piso 1: 15 personas\n• Piso 2: 8 personas\n• Piso 3: 5 personas\n• Zonas especiales',
            'Localizar dónde está cada persona\nDetectar aglomeraciones\nVerificar capacidad de zonas'
        ],
        [
            '📋 TABLA MAIN\n(Área Central)',
            'Centro de pantalla\n(la mayoría del espacio)',
            'Listado detallado de TODOS los accesos hoy:\n• Hora exacta\n• Nombre de persona\n• Documento\n• Tipo (Entrada/Salida)\n• Estado (Dentro/Fuera)\n• Foto/Botones',
            'Ver cada movimiento en detalle\nBuscar persona específica\nHacer clic para más información'
        ],
        [
            '🔍 FILTROS\n(Arriba de tabla)',
            'Superior de la tabla\ncontroles horizontales',
            '• Tipo: ENTRADA/SALIDA/TODOS\n• Búsqueda: por nombre\n• Búsqueda: por documento\n• Rango de fechas',
            'Filtrar información específica\nEncontrar a personas\nAgilizar búsquedas'
        ],
    ]
    
    create_big_table(doc, componentes_data, '1a5200', [1.3, 1.5, 2.5, 2.2])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Monitor en Tiempo Real
    add_heading_styled(doc, 'Monitor en Tiempo Real: Cómo Funciona', 2)
    
    realtime_data = [
        ['CARACTERÍSTICA', 'QUÉ OCURRE', 'POR QUÉ ES IMPORTANTE'],
        [
            'Fila Nueva se Destaca',
            'Cada vez que alguien escanea, aparece nueva fila. Se ilumina con fondo VERDE por 3 segundos indicando registro NUEVO.',
            'Vigilante ve inmediatamente quién acaba de llegar. No hay que buscar.'
        ],
        [
            'Color Verde = Entrada',
            'Las filas con personaas que ENTRARON aparecen con fondo verde o ícono de flecha hacia adentro.',
            'Identificar rápidamente quién está llegando. Distinguir entradas de salidas.'
        ],
        [
            'Color Naranja = Salida',
            'Las filas con personas que SALIERON muestran fondo naranja o ícono de flecha hacia afuera.',
            'Saber quién se va. Controlar que todos salgan al final del día.'
        ],
        [
            'Ordenamiento Automático',
            'La tabla siempre muestra primero los registros más recientes (últimas personas que escanearon).',
            'Vigilante no cansa de scroll. Los nuevos eventos siempre están arriba visible.'
        ],
        [
            'Auto-Actualización',
            'Cada 2 segundos, la tabla se refresca automáticamente sin que hagas nada.',
            'Información siempre actualizada. No necesitas hacer refresh manual (F5).'
        ],
        [
            'Click para Detalles',
            'Haciendo clic en una fila, aparece modal con:\n• Foto grande\n• Programa/especialidad\n• Email y teléfono\n• Historial completo de esa persona',
            'Obtener información profunda rápidamente. Investigar historial de alguien.'
        ],
    ]
    
    create_big_table(doc, realtime_data, '1a5200', [1.8, 2.5, 2.2])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Alertas
    add_heading_styled(doc, 'Sistema de Alertas: Detectar Problemas', 2)
    
    alertas_data = [
        ['TIPO ALERTA', 'QUÉ SIGNIFICA', 'ACCIÓN QUE TOMAR'],
        [
            '🚨 Acceso\nDenegado',
            'Una persona intentó ingresar pero fue RECHAZADA por el sistema. Motivo: no está autorizada, cuenta suspendida, o vencimiento de autorización.',
            'Interrogar a la persona. Verificar identidad. Contactar con administración si insiste. Anotarnotación de incidente.'
        ],
        [
            '⚠️ Documento\nNo Encontrado',
            'Alguien escanió su código pero NO aparece en el sistema. No está registrada aún o hay error en BD.',
            'Acercarse a la persona. Tomar datos manualmente. Registrar como visitante temporal. Avisar a administración.'
        ],
        [
            '⏰ Acceso\nFuera de Horario',
            'Una persona entró fuera de los horarios permitidos (antes de apertura o después de cierre). Puede ser personal autorizado o error de configuración.',
            'Verificar si es instructor, administrador o personal. Revisar configuración de horarios. Reportar si es anómalo.'
        ],
        [
            '👥 Concentración\nAnormal',
            'Un piso excede la capacidad máxima permitida. Hay demasiadas personas en una zona.',
            'Hacer ronda. Verificar que no hay peligro de aglomeración. Comunicar a administración. Limitar entrada si es necesario.'
        ],
        [
            '🚪 Puerta Abierta\nAnormalmente',
            'Si hay cerradura electrónica (ESP32), se alerta si puerta permanece abierta demasiado tiempo.',
            'Verificar que puerta esté segura. Cerrar si es necesario. Revisar si hay falla de dispositivo.'
        ],
        [
            '❌ Falta de\nSalida',
            'Una persona que entró no ha salido después de muchas horas (ej: 8+ horas). Posible olvido o problema.',
            'Contactar a la persona por teléfono o ir a verificar. Asegurarse de que esté bien. Registrar salida si se fue.'
        ],
    ]
    
    create_big_table(doc, alertas_data, '1a5200', [1.5, 2.5, 2.5])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Tareas Diarias Vigilante
    add_heading_styled(doc, 'Tareas Diarias del Vigilante', 2)
    
    tareas_vig_data = [
        ['PERÍODO', 'HORAS', 'TAREAS', 'COSAS A REVISAR'],
        [
            '🌅 MAÑANA\n(Apertura)',
            '6:00 -\n8:30',
            '• Llegar temprano\n• Abrir el sistema\n• Verificar dispositivos\n• Recibir personas',
            '✓ Lector funciona?\n✓ Pantalla enciende?\n✓ Internet activo?\n✓ Hora correcta?'
        ],
        [
            '☀️ DURANTE DÍA\n(Monitoreo)',
            '8:30 -\n12:00',
            '• Vigilancia continua\n• Monitorear panel\n• Revisar alertas\n• Responder consultas',
            '✓ Nuevos accesos normales?\n✓ Hay alertas?\n✓ Personas no autorizadas?\n✓ Aglomeraciones?'
        ],
        [
            '🍽️ ALMUERZO\n(Zona Baja)',
            '12:00 -\n1:00',
            '• Vigilancia reducida\n• Monitoreo de alertas\n• Hacer rondas\n• Control visual',
            '✓ Movimientos anormales?\n✓ Puertas seguras?\n✓ Falta alguien salir?\n✓ Visitantes irregulares?'
        ],
        [
            '🌤️ TARDE\n(Salidas)',
            '1:00 -\n5:00',
            '• Control de salidas\n• Vigilancia de zona\n• Preparar reporte\n• Alertar tardíos',
            '✓ Todos saliendo en orden?\n✓ Hay curiosidades?\n• Personas retenidas?\n✓ Evento especial?'
        ],
        [
            '🌆 CIERRE\n(Finalización)',
            '5:00 -\n6:00',
            '• Verificar todos salieron\n• Generar reporte\n• Documentar eventos\n• Cerrar sesión',
            '✓ Edificio vacío?\n✓ Puertas cerradas?\n✓ Luces apagadas?\n✓ Incidentes registrados?'
        ],
    ]
    
    create_big_table(doc, tareas_vig_data, '1a5200', [1.2, 1.0, 2.0, 2.3])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # ===============================================
    # MÓDULO 3: ADMINISTRACIÓN
    # ===============================================
    
    add_heading_styled(doc, 'MÓDULO 3: PANEL DE ADMINISTRACIÓN', 1)
    
    add_paragraph_justified(doc,
        'El Panel de Administración es el centro de control total. Solo administradores pueden acceder. '
        'Gestión de usuarios, registros, reportes, configuración y auditoría.')
    
    doc.add_paragraph()
    
    # Tabla: Tipos de Admin
    add_heading_styled(doc, 'Tipos de Administradores y Permisos', 2)
    
    admin_tipos_data = [
        ['TIPO', 'ACCESO A...', 'PUEDE HACER', 'NO PUEDE HACER'],
        [
            '👑 SUPERADMIN\n(Principal)',
            '✅ TODO sin restricción\n✅ Logs completos\n✅ Configuración\n✅ Respaldos',
            '• Crear/editar/eliminar usuarios\n• Ver auditoría completa\n• Cambiar cualquier configuración\n• Restaurar respaldos\n• Eliminar datos permanentemente',
            '❌ Nada - acceso total'
        ],
        [
            '🛡️ ADMIN\nNORMAL',
            '✅ Gestión usuarios\n✅ Reportes\n✅ Algunos parámetros\n⚠️ Auditoría limitada',
            '• Crear usuarios (no eliminar)\n• Editar algunos datos\n• Ver reportes estándar\n• Cambiar parámetros básicos\n• Registrar personas',
            '❌ Eliminar usuarios\n❌ Ver auditoría completa\n❌ Restaurar respaldos\n❌ Cambios críticos'
        ],
    ]
    
    create_big_table(doc, admin_tipos_data, '1a5200', [1.3, 1.8, 2.0, 1.9])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Gestión de Usuarios
    add_heading_styled(doc, 'Gestión de Usuarios (Crear y Editar Cuentas)', 2)
    
    gestion_usuarios_data = [
        ['ACCIÓN', 'PASOS', 'CAMPOS A LLENAR', 'RESULTADO'],
        [
            '📝 Crear\nNuevo\nUsuario',
            '1. Click "Crear Usuario"\n2. Llenar formulario\n3. Click "GUARDAR"\n4. Sistema envía email',
            '• Email (único)\n• Contraseña (fuerte)\n• Nombre completo\n• Rol (Vigilante/Admin)\n• Activo: Sí/No',
            'Usuario creado\nRecibe email\nPuede iniciar sesión'
        ],
        [
            '✏️ Editar\nUsuario',
            '1. Buscar usuario en lista\n2. Click "EDITAR"\n3. Modificar campos\n4. Click "GUARDAR"',
            '• Nombre\n• Email\n• Rol\n• Estado activo/inactivo\n(NO contraseña aquí)',
            'Cambios aplicados\nSe registra en auditoría\nEmail notificado'
        ],
        [
            '🔑 Resetear\nContraseña',
            '1. Seleccionar usuario\n2. Click "RESETEAR CONTRASEÑA"\n3. Sistema envía email\n4. Usuario recupera contraseña',
            '→ Sistema envía enlace al email del usuario\n→ Usuario crea nueva contraseña',
            'Email enviado\nUsuario recupera acceso\nContraseña anterior anulada'
        ],
        [
            '🔴 Desactivar\nCuenta',
            '1. Click usuario\n2. Cambiar "Activo" → NO\n3. Click "GUARDAR"\n4. ⚠️ No eliminar, solo desactivar',
            '• Cambiar estado a INACTIVO\n• Guardar cambio\n• Opcional: comentario de por qué',
            'Usuario NO puede loguear\nCuenta preservada (auditoría)\nPuede reactivarse'
        ],
        [
            '🗑️ Eliminar\nUsuario\n(SuperAdmin)',
            '1. Click usuario\n2. Click "ELIMINAR"\n3. Confirmar diálogo\n4. ⚠️ PERMANENTE',
            '→ SuperAdmin only\n→ Requiere confirmación\n→ No se puede deshacer',
            'Usuario eliminado permanentemente\nNo se puede recuperar\nRegistrado en auditoría'
        ],
    ]
    
    create_big_table(doc, gestion_usuarios_data, '1a5200', [1.2, 1.8, 2.0, 1.5])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Registro de Personas
    add_heading_styled(doc, 'Registro de Personas: Tipos y Procedimiento', 2)
    
    personas_tipos_data = [
        ['TIPO', 'DESCRIPCIÓN', 'DATOS REQUERIDOS', 'ESPECIAL'],
        [
            '👨‍🎓 APRENDIZ',
            'Estudiante en programa de formación',
            '• Nombre\n• Documento\n• Email, teléfono\n• Programa\n• Ficha\n• Especialidad',
            '✓ Es el 90% de registros\n✓ Puede escanear código\n✓ Ver historial completo'
        ],
        [
            '👨‍🏫 INSTRUCTOR',
            'Personal docente y formador',
            '• Nombre\n• Documento\n• Email, teléfono\n• Departamento\n• Especialidad\n• Horario (opcional)',
            '✓ Acceso a más áreas\n✓ Horarios flexibles\n✓ Reportes especiales'
        ],
        [
            '👔 ADMINISTRATIVO',
            'Personal: secretaria, mantenimiento, servicios',
            '• Nombre\n• Documento\n• Email, teléfono\n• Cargo\n• Departamento\n• Horario',
            '✓ Acceso todo edificio\n✓ Horarios extendidos\n✓ Registro diferente'
        ],
        [
            '👨‍💼 VISITANTE',
            'Personas de afuera (proveedor, padre, etc.)',
            '• Nombre\n• Documento\n• Email, teléfono\n• Motivo de visita\n• Fecha inicio\n• Fecha fin',
            '⚠️ Acceso acotado\n⚠️ Tiempo limitado\n⚠️ Auto-bloqueo al vencer fecha'
        ],
    ]
    
    create_big_table(doc, personas_tipos_data, '1a5200', [1.2, 1.8, 2.0, 1.5])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Cómo Registrar Persona
    add_heading_styled(doc, 'Pasos para Registrar una Persona Nueva', 2)
    
    registro_pasos_data = [
        ['PASO', 'ACCIÓN', 'DETALLES', 'IMPORTANTE'],
        [
            '1',
            'Acceder a registro',
            'Dashboard → Click "REGISTRO DE PERSONAS" o "+ REGISTRAR"',
            '✓ Solo admin puede registrar'
        ],
        [
            '2',
            'Datos básicos',
            'Llenar campos obligatorios (*):\n• Nombre completo\n• Tipo documento\n• Número documento\n• Email, teléfono',
            '⚠️ Documento único\n⚠️ Email válido'
        ],
        [
            '3',
            'Seleccionar perfil',
            'Dropdown "Perfil":\n- Aprendiz\n- Instructor\n- Administrativo\n- Visitante',
            '✓ Determina permisos\n✓ Según perfil aparecen campos adicionales'
        ],
        [
            '4',
            'Datos específicos',
            'Llenar según perfil elegido:\n• Si Aprendiz: programa, ficha, especialidad\n• Si Visitante: fecha inicio/fin',
            '✓ Varían por tipo\n✓ Algunos campos aparecen/desaparecen'
        ],
        [
            '5',
            'Subir foto',
            'Click "SUBIR FOTO" o arrastra imagen\n• Formatos: JPG, PNG\n• Tamaño recomendado: 200x200px',
            '✓ Opcional pero recomendado\n✓ Aparece en vigilancia'
        ],
        [
            '6',
            'Guardar',
            'Click "GUARDAR" o "REGISTRAR PERSONA"',
            '✅ Persona crea en BD\n✅ Ya puede escanear código\n✅ Email de confirmas (opcional)'
        ],
    ]
    
    create_big_table(doc, registro_pasos_data, '1a5200', [0.6, 1.5, 3.0, 1.4])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Búsqueda y Edición
    add_heading_styled(doc, 'Búsqueda, Edición y Control de Personas', 2)
    
    busqueda_data = [
        ['OPERACIÓN', 'CÓMO HACERLO', 'RESULTADO', 'AUDITORÍA'],
        [
            '🔍 Buscar persona',
            'Área "BÚSQUEDA":\n• Por nombre (completo o parcial)\n• Por documento\n• Por email',
            'Resultados instantáneos mostrando:\n• Nombre\n• Documento\n• Perfil\n• Estado (activo/inactivo)',
            '✓ Búsqueda registrada'
        ],
        [
            '📋 Ver detalles',
            'Click en resultado\nSe abre overlay con info completa',
            'Mostrará:\n• Datos personales\n• Foto\n• Historial de accesos (entrada/salida)\n• Cambios previos',
            '✓ Visualización registrada'
        ],
        [
            '✏️ Editar información',
            '1. Buscar persona\n2. Click "EDITAR"\n3. Cambiar campos\n4. "GUARDAR CAMBIOS"',
            'Campos editables:\n• Nombre, email, teléfono\n• Programa/especialidad\n• Foto\n• Estado activo/inactivo',
            '✓ Cambio registrado\n✓ Muestra quién y cuándo\n✓ Valor anterior guardado'
        ],
        [
            '🚫 Desactivar acceso',
            '1. Editar persona\n2. Cambiar "Activo" → NO\n3. Guardar',
            'Resultado:\n• Persona no puede escanear\n• Sistema rechaza su código\n• Cuenta existe pero bloqueada',
            '✓ Registro completo\n✓ Se puede reactivar\n✓ No se elimina'
        ],
        [
            '🔑 Visitante: Autorizar',
            'Registrar con Perfil VISITANTE\nEstablecer fechas inicio y fin',
            'El sistema:\n• Permite acceso hasta fecha fin\n• Después de la fecha: acceso automáticamente denegado\n• No requiere intervención admin',
            '✓ Automático\n✓ Sin intervención\n✓ Registrado en logs'
        ],
    ]
    
    create_big_table(doc, busqueda_data, '1a5200', [1.3, 2.0, 2.2, 1.5])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Reportes
    add_heading_styled(doc, 'Generación de Reportes: Tipos y Uso', 2)
    
    reportes_data = [
        ['TIPO DE REPORTE', 'QUÉ CONTIENE', 'ÚTIL PARA', 'FORMATO'],
        [
            '📅 ASISTENCIA\nDIARIA',
            '• Quién entró/salió\n• Hora entrada/salida\n• Total horas\n• Estado presencia',
            '✓ Control de asistencia\n✓ Cálculo horas\n✓ Justificar faltas',
            'Excel, PDF'
        ],
        [
            '📊 RESUMEN\nMENSUAL',
            '• Mes completo\n• Asistencia\n• Faltas, retrasos\n• Salidas temprano',
            '✓ Análisis mensual\n✓ Recursos\n✓ Ocupación promedio',
            'Excel, PDF'
        ],
        [
            '👥 PERSONAS\nREGISTRADAS',
            '• Lista completa\n• Nombre, documento\n• Perfil, programa\n• Email, teléfono',
            '✓ Control de población\n✓ Auditoría general\n✓ Comunicaciones masivas',
            'Excel'
        ],
        [
            '📍 POR ZONA\n/PISO',
            '• Accesos por zona\n• Frecuencia\n• Horarios\n• Patrones',
            '✓ Seguridad por ubicación\n✓ Ocupación\n✓ Optimizar espacios',
            'Excel'
        ],
        [
            '⚠️ ALERTAS\nY ANOMALÍAS',
            '• Accesos denegados\n• Documentos no encontrados\n• Fallos seguridad\n• Incidentes',
            '✓ Investigación\n✓ Seguridad\n✓ Reporte directivas',
            'Excel, PDF'
        ],
        [
            '💰 INGRESOS\n(si aplica)',
            '• Pagos por servicios\n• Cafetería, fotocopia\n• Totales por día/mes\n• Por usuario',
            '✓ Finanzas\n✓ Auditoría contable\n✓ Análisis económico',
            'Excel'
        ],
    ]
    
    create_big_table(doc, reportes_data, '1a5200', [1.5, 2.0, 1.8, 1.2])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Descargar Reportes
    add_heading_styled(doc, 'Cómo Generar y Descargar Reportes', 2)
    
    download_data = [
        ['PASO', 'ACCIÓN', 'DETALLES'],
        ['1', 'Click en REPORTES', 'Menú admin → Opción "REPORTES" o "GENERAR REPORTE"'],
        ['2', 'Elegir tipo', 'Select dropdown: Asistencia, Mensual, Personas, Zonas, Alertas, Ingresos'],
        ['3', 'Rango de fechas', 'Desde [fecha] Hasta [fecha]. Ej: 01/04/2026 al 30/04/2026'],
        ['4', 'Filtros extras', 'Si quieres: persona específica, zona, tipo evento, etc.'],
        ['5', 'Click GENERAR', 'Sistema procesa... (5-30 segundos según tamaño)'],
        ['6', 'Esperar completa', 'Cuando termina, aparece botón "DESCARGAR"'],
        ['7', 'Descargar', 'Click "DESCARGAR EXCEL" o "DESCARGAR PDF"'],
        ['8', 'Abrir archivo', 'Se descarga. Abre en Excel / Acrobat. Puedes editar, gráficos, imprimir.'],
    ]
    
    create_big_table(doc, download_data, '1a5200', [0.8, 1.8, 3.9])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Configuración
    add_heading_styled(doc, 'Configuración del Sistema', 2)
    
    config_data = [
        ['PARÁMETRO', 'QUÉ ES', 'CÓMO CAMBIARLO', 'IMPACTO'],
        [
            '🌐 IP del ESP32',
            'Dirección física del dispositivo cerradura electrónica',
            'Campo texto: Ingresar "192.168.1.45"\nBotón "Probar conexión"\nClick "Guardar"',
            'Sistema se conecta a puerta'
        ],
        [
            '🏢 Pisos/Zonas',
            'Cuántos pisos tiene edificio y sus nombres',
            '1. Click "Agregar piso"\n2. Nombre: "Piso 1"\n3. Capacidad máxima: "50"\n4. Guardar',
            'Aparecen en panel vigilancia'
        ],
        [
            '🕐 Horarios\nPermitidos',
            'A qué horas permite entrada/salida',
            'Hora apertura: 6:00 AM\nHora cierre: 6:00 PM\nAcceso fuera horario: ¿Permitir? Sí/No',
            'Alertas si acceso fuera de horario'
        ],
        [
            '📢 Mensajes\nPersonalizados',
            'Textos que ven aprendices en pantalla',
            'Campo texto editable:\n"Bienvenido al SENA"\n"Error: Intenta de nuevo"',
            'Experiencia personalizada'
        ],
        [
            '🔐 Seguridad\n(SuperAdmin)',
            'Nivel de encriptación, días de retención logs',
            'Opciones avanzadas:\n• Encriptación: ON/OFF\n• Retención logs: 30 días\n• Respaldos: Automático/Manual',
            'Protección de datos'
        ],
    ]
    
    create_big_table(doc, config_data, '1a5200', [1.3, 1.5, 2.2, 1.5])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Tareas Críticas Admin
    add_heading_styled(doc, 'Tareas Críticas del Administrador', 2)
    
    tareas_admin_data = [
        ['FRECUENCIA', 'TAREAS', 'VERIFICACIONES', 'SI ALGO FALLA'],
        [
            '🌅 DIARIO\n(Mañana)',
            '✓ Verificar sistema UP\n✓ Probar lector\n✓ Revisar alertas noche',
            '• ¿Pantalla enciende?\n• ¿Lector funciona?\n• ¿Internet OK?\n• ¿Hay alertas pendientes?',
            'Contactar técnico si falla'
        ],
        [
            '📅 SEMANAL',
            '✓ Revisar nuevas personas\n✓ Desactivar visitantes vencidos\n✓ Auditoría cambios\n✓ Contesta consultas',
            '• ¿Visitantes autorizados?\n• ¿Accesos anormales?\n• ¿Cambios no autorizados?\n• ¿Documentos faltantes?',
            'Investigar anomalía'
        ],
        [
            '📊 MENSUAL',
            '✓ Generar reportes\n✓ Análisis estadísticas\n✓ Hacer respaldo datos\n✓ Revisar auditoría mes',
            '• ¿Ocupación normal?\n• ¿Alertas patrones?\n• ¿Respaldo completo?\n• ¿Cambios registrados?',
            'Revisar y documentar'
        ],
        [
            '🔐 TRIMESTRAL',
            '✓ Cambiar psw SuperAdmin\n✓ Revisar permisos usuarios\n✓ Verificar respaldos\n✓ Reportear directivas',
            '• ¿Usuarios tienen acceso correcto?\n• ¿Respaldos funcionan?\n• ¿Estadísticas OK?\n• ¿Seguridad intacta?',
            'Resolver antes del siguiente'
        ],
        [
            '🚨 EVENTOS',
            '✓ Acceso denegado: investigar\n✓ Fallo sistema: reiniciar\n✓ Incidente seguridad: alertar\n✓ Visitante: verificar',
            '• ¿Quién es la persona?\n• ¿Hay log del evento?\n• ¿Dirección está enterada?\n• ¿Cambio PSW necesario?',
            'Documentar y seguimiento'
        ],
    ]
    
    create_big_table(doc, tareas_admin_data, '1a5200', [1.0, 2.0, 2.0, 1.5])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # ==== CONCLUSIÓN ====
    add_heading_styled(doc, 'Conclusión: Resumen de los Tres Módulos', 1)
    
    resumen_data = [
        ['MÓDULO', 'USUARIOS', 'AL DÍA HACEN...', 'HERRAMIENTA PRINCIPAL'],
        [
            '🎯 INGRESO',
            'Aprendices\nInstructores\nVisitantes',
            'Escanean código de barras\nRegistran entrada/salida\nReciben confirmación',
            'Lector código de barras\nPantalla de ingreso'
        ],
        [
            '👮 VIGILANCIA',
            'Vigilantes\nSeguridad',
            'Monitorean en tiempo real\nDetectan alertas\nControlan acceso',
            'Panel tiempo real\nFiltros y búsqueda\nControl puertas'
        ],
        [
            '🛡️ ADMINISTRACIÓN',
            'Administradores\nSuperAdmin',
            'Registran personas\nGestionan usuarios\nGeneran reportes\nConfiguran sistema',
            'Base de datos\nFunciones avanzadas\nAuditoría completa'
        ],
    ]
    
    create_big_table(doc, resumen_data, '1a5200', [1.5, 1.8, 2.2, 2.0])
    doc.add_paragraph()
    doc.add_paragraph()
    
    conclusion = doc.add_paragraph(
        'Los tres módulos trabajanen sinergia: Ingreso registra, Vigilancia monitorea, Administración gestiona. '
        'Juntos forman un sistema integrado de seguridad y control completo.')
    conclusion.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in conclusion.runs:
        run.font.size = Pt(12)
        run.italic = True
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    footer = doc.add_paragraph('Sistema de Control de Ingresos SENA © 2026')
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(120, 120, 120)
    
    # ==== GUARDAR ====
    output_path = r'C:\Users\AMAT\Desktop\ingreso de aprendices\INFORME_SISTEMA_CON_TABLAS.docx'
    doc.save(output_path)
    print(f'✅ Documento guardado en: {output_path}')
    print(f'📄 Archivo: INFORME_SISTEMA_CON_TABLAS.docx')
    print(f'📊 Estructura: INGRESO CON TABLAS > VIGILANCIA CON TABLAS > ADMINISTRACIÓN CON TABLAS')

if __name__ == '__main__':
    main()
