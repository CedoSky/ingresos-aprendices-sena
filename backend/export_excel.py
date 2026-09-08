"""
Exportación profesional a Excel — Sistema de Control de Acceso SENA
Genera tablas completas con toda la información de accesos y equipos.
Guarda automáticamente en USB detectada y en RESPALDOS_USB local.
"""

import os
import io
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

# ─── Paleta institucional SENA ────────────────────────────────────────────────
AZ1 = "003087"; AZ2 = "005EB8"; AZ3 = "00AEEF"
VE1 = "1A5E31"; VE2 = "D6F0DC"
RO1 = "C00000"; RO2 = "FFE0E0"
AM1 = "7F6000"; AM2 = "FFF2CC"
MO1 = "4B0082"; MO2 = "F0E6FF"
BL  = "FFFFFF"; GR1 = "F2F2F2"; GR2 = "595959"; GR3 = "D9D9D9"
PA1 = "EBF3FB"   # fila par


# ─── Helpers ─────────────────────────────────────────────────────────────────
def _f(c):
    return PatternFill("solid", fgColor=c)

def _ft(bold=False, color="000000", size=10, italic=False):
    return Font(bold=bold, color=color, size=size, name="Calibri", italic=italic)

def _ac():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _al():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def _ar():
    return Alignment(horizontal="right", vertical="center", wrap_text=False)

def _borde(color=GR3):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def _borde_cab():
    g = Side(style="medium", color=AZ1)
    t = Side(style="thin",   color=GR3)
    return Border(left=t, right=t, top=g, bottom=g)

def _mc(ws, r1, c1, r2, c2):
    """merge_cells con keyword args — sintaxis correcta de openpyxl."""
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


# ─── Cabecera institucional ───────────────────────────────────────────────────
def _cabecera(ws, titulo, n, fecha_gen):
    """Inserta encabezado en filas 1-6. Devuelve fila 7 (primera libre)."""

    # Barra superior
    _mc(ws, 1, 1, 1, n)
    ws.cell(row=1, column=1).fill = _f(AZ1)
    ws.row_dimensions[1].height = 7

    # Bloque logo SENA (cols 1-3, filas 2-4)
    for r in range(2, 5):
        for c in range(1, 4):
            ws.cell(row=r, column=c).fill = _f(AZ1)
    _mc(ws, 2, 1, 4, 3)
    c = ws.cell(row=2, column=1)
    c.value = "SENA"; c.fill = _f(AZ1)
    c.font = Font(bold=True, color=BL, size=30, name="Calibri")
    c.alignment = _ac()

    # Título (cols 4-n, filas 2-3)
    _mc(ws, 2, 4, 3, n)
    c = ws.cell(row=2, column=4)
    c.value = titulo.upper(); c.fill = _f(AZ1)
    c.font = Font(bold=True, color=BL, size=14, name="Calibri")
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

    # Subtítulo (fila 4, cols 4-n)
    _mc(ws, 4, 4, 4, n)
    c = ws.cell(row=4, column=4)
    c.value = "Servicio Nacional de Aprendizaje  ·  Sistema de Control de Acceso"
    c.fill = _f(AZ2)
    c.font = Font(italic=True, color=BL, size=10, name="Calibri")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 18

    # Fila 5: metadata
    mc = max(n // 3, 1)
    ws.row_dimensions[5].height = 17

    _mc(ws, 5, 1, 5, mc)
    c = ws.cell(row=5, column=1)
    c.value = f"  Generado: {fecha_gen}"
    c.fill = _f(GR1); c.font = _ft(bold=True, color=GR2, size=9); c.alignment = _al()

    _mc(ws, 5, mc + 1, 5, mc * 2)
    c = ws.cell(row=5, column=mc + 1)
    c.value = "Documento confidencial · Uso interno SENA"
    c.fill = _f(GR1); c.font = _ft(italic=True, color=GR2, size=9); c.alignment = _ac()

    _mc(ws, 5, mc * 2 + 1, 5, n)
    ws.cell(row=5, column=mc * 2 + 1).fill = _f(GR1)

    # Línea separadora (fila 6)
    _mc(ws, 6, 1, 6, n)
    ws.cell(row=6, column=1).fill = _f(AZ3)
    ws.row_dimensions[6].height = 4

    return 7   # primera fila de datos


def _header_row(ws, fila, cols, anchos):
    ws.row_dimensions[fila].height = 28
    for i, (txt, ancho) in enumerate(zip(cols, anchos), 1):
        c = ws.cell(row=fila, column=i, value=txt)
        c.fill = _f(AZ2)
        c.font = _ft(bold=True, color=BL, size=10)
        c.alignment = _ac()
        c.border = _borde_cab()
        ws.column_dimensions[get_column_letter(i)].width = ancho
    return fila + 1


def _data_row(ws, fila, valores, par, altura=18):
    bg = PA1 if par else BL
    ws.row_dimensions[fila].height = altura
    for i, v in enumerate(valores, 1):
        c = ws.cell(row=fila, column=i, value=v)
        c.fill = _f(bg); c.font = _ft(size=10)
        c.alignment = _al(); c.border = _borde()
    return bg


def _pie(ws, fila, n, texto):
    ws.row_dimensions[fila].height = 15
    _mc(ws, fila, 1, fila, n)
    c = ws.cell(row=fila, column=1, value=texto)
    c.fill = _f(GR1); c.font = _ft(italic=True, color=GR2, size=9)
    c.alignment = _ar()
    c.border = Border(top=Side(style="medium", color=AZ2))


def _color_movimiento(ws, fila, col, tipo):
    c = ws.cell(row=fila, column=col)
    c.alignment = _ac()
    if tipo == "ENTRADA":
        c.fill = _f(VE2); c.font = _ft(bold=True, color=VE1, size=10)
    else:
        c.fill = _f(RO2); c.font = _ft(bold=True, color=RO1, size=10)


def _color_perfil(ws, fila, col, perfil):
    c = ws.cell(row=fila, column=col)
    c.alignment = _ac()
    pf = perfil.upper()
    if pf == "APRENDIZ":
        c.fill = _f("D6E4FF"); c.font = _ft(bold=True, color=AZ1, size=10)
    elif pf == "INSTRUCTOR":
        c.fill = _f(AM2); c.font = _ft(bold=True, color=AM1, size=10)
    elif pf == "VISITANTE":
        c.fill = _f(MO2); c.font = _ft(bold=True, color=MO1, size=10)


# ═══════════════════════════════════════════════════════════════════════════════
# HOJA 1 — DIARIO COMPLETO DE ACCESOS
# ═══════════════════════════════════════════════════════════════════════════════
def hoja_diario(wb, registros, fecha_gen):
    ws = wb.create_sheet("Diario de Accesos", 0)
    ws.sheet_view.showGridLines = False

    COLS = [
        ("N°",4), ("FECHA",13), ("HORA",10),
        ("NOMBRE COMPLETO",36), ("N° DOCUMENTO",18), ("TIPO DOC",11),
        ("PERFIL",13), ("MOVIMIENTO",14), ("AMBIENTE / LUGAR",24),
        ("REGISTRÓ EQUIPO",16), ("CÓDIGO EQUIPO",20), ("MARCA / MODELO",22),
        ("PROGRAMA",28), ("FICHA",11),
    ]
    cab  = [c[0] for c in COLS]
    anch = [c[1] for c in COLS]
    n    = len(COLS)

    fila = _cabecera(ws, "Diario Completo de Ingresos y Salidas", n, fecha_gen)
    fila = _header_row(ws, fila, cab, anch)

    entradas = salidas = con_eq = 0
    for i, r in enumerate(registros, 1):
        par = (i % 2 == 0)
        _data_row(ws, fila, [
            i, r["fecha"], r["hora"],
            r["nombre"], r["numero_doc"], r["tipo_doc"], r["perfil"],
            r["tipo"], r["ambiente"],
            "",                  # col 10 — se sobreescribe abajo
            r["eq_codigo"], r["eq_nombre"],
            r["programa"], r["ficha"],
        ], par)

        _color_movimiento(ws, fila, 8, r["tipo"])
        _color_perfil(ws, fila, 7, r["perfil"])

        if r["tipo"] == "ENTRADA": entradas += 1
        else: salidas += 1

        # REGISTRÓ EQUIPO
        ce = ws.cell(row=fila, column=10); ce.alignment = _ac()
        if r["eq_codigo"]:
            ce.value = "✔  SÍ"
            ce.fill = _f(VE2); ce.font = _ft(bold=True, color=VE1, size=10)
            con_eq += 1
        else:
            ce.value = "NO"; ce.font = _ft(color="AAAAAA", size=10)

        ws.cell(row=fila, column=11).alignment = _ac()
        ws.cell(row=fila, column=12).alignment = _al()
        fila += 1

    _pie(ws, fila, n,
         f"  Total: {len(registros)}  ·  Entradas: {entradas}  ·  "
         f"Salidas: {salidas}  ·  Con equipo: {con_eq}   ")

    ws.freeze_panes = "A8"
    ws.auto_filter.ref = f"A7:{get_column_letter(n)}{fila - 1}"


# ═══════════════════════════════════════════════════════════════════════════════
# HOJA 2 — RESUMEN DEL DÍA  (una fila por persona con entrada + salida)
# ═══════════════════════════════════════════════════════════════════════════════
def hoja_resumen_dia(wb, registros, stats, fecha_gen):
    ws = wb.create_sheet("Resumen del Día")
    ws.sheet_view.showGridLines = False

    hoy = datetime.now().strftime("%d/%m/%Y")
    hoy_reg = [r for r in registros if r["fecha"] == hoy]

    COLS = [
        ("N°",4), ("NOMBRE COMPLETO",36), ("N° DOCUMENTO",18),
        ("PERFIL",13), ("HORA ENTRADA",14), ("HORA SALIDA",14),
        ("AMBIENTE",24), ("REGISTRÓ EQUIPO",16),
        ("CÓDIGO EQUIPO",20), ("PROGRAMA",28), ("FICHA",11),
    ]
    cab  = [c[0] for c in COLS]
    anch = [c[1] for c in COLS]
    n    = len(COLS)

    fila = _cabecera(ws, f"Resumen de Actividad del Día  ·  {hoy}", n, fecha_gen)

    # Tabla de estadísticas del día (4 celdas en una fila)
    _mc(ws, fila, 1, fila, n)
    c = ws.cell(row=fila, column=1, value="  ESTADÍSTICAS DEL DÍA")
    c.fill = _f(AZ1); c.font = _ft(bold=True, color=BL, size=11); c.alignment = _al()
    ws.row_dimensions[fila].height = 22
    fila += 1

    stats_items = [
        ("Total entradas hoy",    stats["entradas_hoy"]),
        ("Total salidas hoy",     stats["salidas_hoy"]),
        ("Personas con equipo",   stats.get("con_equipo_hoy", 0)),
        ("Total personas sistema",stats["total_personas"]),
    ]
    ws.row_dimensions[fila].height = 22
    col_offset = 1
    for lbl, val in stats_items:
        if col_offset + 1 > n:
            break
        cl = ws.cell(row=fila, column=col_offset, value=lbl)
        cl.fill = _f(GR1); cl.font = _ft(color=GR2, size=10)
        cl.alignment = _al(); cl.border = _borde()
        cv = ws.cell(row=fila, column=col_offset + 1, value=val)
        cv.fill = _f(PA1); cv.font = _ft(bold=True, color=AZ1, size=13)
        cv.alignment = _ac(); cv.border = _borde()
        col_offset += 3
    fila += 2

    # Encabezado de la tabla del día
    _mc(ws, fila, 1, fila, n)
    c = ws.cell(row=fila, column=1, value=f"  MOVIMIENTOS DEL DÍA — {hoy}  ({len(hoy_reg)} registros)")
    c.fill = _f(AZ2); c.font = _ft(bold=True, color=BL, size=11); c.alignment = _al()
    ws.row_dimensions[fila].height = 22
    fila += 1

    fila = _header_row(ws, fila, cab, anch)

    # Agrupar por persona: una fila con hora_entrada + hora_salida
    personas_dia: dict = {}
    for r in sorted(hoy_reg, key=lambda x: x["hora"]):
        k = r["numero_doc"]
        if k not in personas_dia:
            personas_dia[k] = {
                "nombre": r["nombre"], "numero_doc": r["numero_doc"],
                "perfil": r["perfil"], "programa": r["programa"],
                "ficha": r["ficha"], "ambiente": r["ambiente"],
                "eq_codigo": r["eq_codigo"], "eq_nombre": r["eq_nombre"],
                "hora_entrada": "", "hora_salida": "",
            }
        if r["tipo"] == "ENTRADA":
            personas_dia[k]["hora_entrada"] = r["hora"]
            personas_dia[k]["ambiente"]  = r["ambiente"]
            personas_dia[k]["eq_codigo"] = r["eq_codigo"]
            personas_dia[k]["eq_nombre"] = r["eq_nombre"]
        else:
            personas_dia[k]["hora_salida"] = r["hora"]

    for i, p in enumerate(personas_dia.values(), 1):
        par = (i % 2 == 0)
        _data_row(ws, fila, [
            i, p["nombre"], p["numero_doc"], p["perfil"],
            p["hora_entrada"], p["hora_salida"],
            p["ambiente"],
            "✔  SÍ" if p["eq_codigo"] else "NO",
            p["eq_codigo"], p["programa"], p["ficha"],
        ], par)

        _color_perfil(ws, fila, 4, p["perfil"])

        ce = ws.cell(row=fila, column=8); ce.alignment = _ac()
        if p["eq_codigo"]:
            ce.fill = _f(VE2); ce.font = _ft(bold=True, color=VE1, size=10)
        else:
            ce.font = _ft(color="AAAAAA", size=10)

        fila += 1

    _pie(ws, fila, n, f"  Personas distintas con movimiento hoy: {len(personas_dia)}   ")
    ws.freeze_panes = "A8"


# ═══════════════════════════════════════════════════════════════════════════════
# HOJA 3 — PERSONAS REGISTRADAS
# ═══════════════════════════════════════════════════════════════════════════════
def hoja_personas(wb, personas, fecha_gen):
    ws = wb.create_sheet("Personas Registradas")
    ws.sheet_view.showGridLines = False

    COLS = [
        ("N°",4), ("NOMBRE COMPLETO",36), ("TIPO DOC",11), ("N° DOCUMENTO",20),
        ("PERFIL",13), ("PROGRAMA",28), ("FICHA",12), ("ESPECIALIDAD",24),
        ("ÁREA",20), ("EMAIL",30), ("TELÉFONO",14),
        ("VERIFICADO",12), ("FECHA REGISTRO",18),
    ]
    cab  = [c[0] for c in COLS]
    anch = [c[1] for c in COLS]
    n    = len(COLS)

    fila = _cabecera(ws, "Directorio de Personas Registradas", n, fecha_gen)
    fila = _header_row(ws, fila, cab, anch)

    for i, p in enumerate(personas, 1):
        fr = p.get("created_at", "")
        if fr and "T" in str(fr):
            try:
                fr = datetime.fromisoformat(str(fr)).strftime("%d/%m/%Y")
            except Exception:
                pass

        _data_row(ws, fila, [
            i, p.get("nombre",""), p.get("tipo_doc",""), p.get("numero_doc",""),
            p.get("perfil",""), p.get("programa",""), p.get("ficha",""),
            p.get("especialidad",""), p.get("area",""),
            p.get("email",""), p.get("telefono",""),
            "SÍ" if p.get("verificado") else "NO", fr,
        ], i % 2 == 0)

        _color_perfil(ws, fila, 5, p.get("perfil",""))

        cv = ws.cell(row=fila, column=12); cv.alignment = _ac()
        if p.get("verificado"):
            cv.fill = _f(VE2); cv.font = _ft(bold=True, color=VE1, size=10)
        else:
            cv.font = _ft(color="AAAAAA", size=10)

        fila += 1

    _pie(ws, fila, n, f"  Total personas registradas: {len(personas)}   ")
    ws.freeze_panes = "A8"
    ws.auto_filter.ref = f"A7:{get_column_letter(n)}{fila - 1}"


# ═══════════════════════════════════════════════════════════════════════════════
# HOJA 4 — INVENTARIO DE EQUIPOS
# ═══════════════════════════════════════════════════════════════════════════════
def hoja_equipos(wb, computadores, personas_por_id, fecha_gen):
    ws = wb.create_sheet("Inventario de Equipos")
    ws.sheet_view.showGridLines = False

    COLS = [
        ("N°",4), ("CÓDIGO BARRAS",22), ("N° SERIE",20),
        ("MARCA",16), ("MODELO",20), ("TIPO",14), ("ESTADO",12),
        ("ASIGNADO A",32), ("UBICACIÓN",24), ("FECHA REGISTRO",16),
    ]
    cab  = [c[0] for c in COLS]
    anch = [c[1] for c in COLS]
    n    = len(COLS)

    fila = _cabecera(ws, "Inventario de Equipos Tecnológicos", n, fecha_gen)
    fila = _header_row(ws, fila, cab, anch)

    for i, c in enumerate(computadores, 1):
        fr = c.get("created_at","")
        if fr and "T" in str(fr):
            try:
                fr = datetime.fromisoformat(str(fr)).strftime("%d/%m/%Y")
            except Exception:
                pass

        _data_row(ws, fila, [
            i, c.get("codigo_barras",""), c.get("serie",""),
            c.get("marca",""), c.get("modelo",""), c.get("tipo",""),
            c.get("estado","").capitalize(),
            personas_por_id.get(c.get("asignado_a_id",""), "Sin asignar"),
            c.get("ubicacion",""), fr,
        ], i % 2 == 0)

        ce = ws.cell(row=fila, column=7); ce.alignment = _ac()
        est = c.get("estado","").lower()
        if est == "activo":
            ce.fill = _f(VE2); ce.font = _ft(bold=True, color=VE1, size=10)
        elif est in ("inactivo","baja","dado de baja"):
            ce.fill = _f(RO2); ce.font = _ft(bold=True, color=RO1, size=10)

        fila += 1

    _pie(ws, fila, n, f"  Total equipos: {len(computadores)}   ")
    ws.freeze_panes = "A8"
    ws.auto_filter.ref = f"A7:{get_column_letter(n)}{fila - 1}"


# ─── Detección USB (Windows) ──────────────────────────────────────────────────
def detectar_usb():
    try:
        import ctypes
        bitmask = ctypes.windll.kernel32.GetLogicalDrives()
        for i in range(26):
            if bitmask & (1 << i):
                letra = chr(65 + i)
                if ctypes.windll.kernel32.GetDriveTypeW(f"{letra}:\\") == 2:
                    return f"{letra}:\\"
    except Exception:
        pass
    return None


# ═══════════════════════════════════════════════════════════════════════════════
# FUNCIÓN PRINCIPAL
# ═══════════════════════════════════════════════════════════════════════════════
def exportar_todo_excel(app):
    if not OPENPYXL_OK:
        raise ImportError("openpyxl no instalado. Ejecuta: pip install openpyxl")

    with app.app_context():
        from models import Persona, RegistroAcceso, Computador, RegistroAccesoEquipo

        # ── Consultas ──────────────────────────────────────────────────
        personas_obj     = Persona.query.filter_by(deleted_at=None).order_by(Persona.nombre).all()
        accesos_obj      = (RegistroAcceso.query
                            .filter_by(deleted_at=None)
                            .order_by(RegistroAcceso.timestamp.desc()).all())
        computadores_obj = Computador.query.filter_by(deleted_at=None).order_by(Computador.marca).all()
        eq_accesos_obj   = RegistroAccesoEquipo.query.filter_by(deleted_at=None).all()

        # ── Diccionarios ───────────────────────────────────────────────
        personas_map     = {p.id: p         for p in personas_obj}
        personas_doc_map = {p.numero_doc: p for p in personas_obj}
        personas_por_id  = {p.id: p.nombre  for p in personas_obj}
        equipos_map      = {c.id: c         for c in computadores_obj}

        # índice registro_acceso_id → lista de RegistroAccesoEquipo
        eq_por_acceso: dict = {}
        for eq in eq_accesos_obj:
            if eq.registro_acceso_id:
                eq_por_acceso.setdefault(eq.registro_acceso_id, []).append(eq)

        # fallback: persona_id → lista (para accesos sin registro_acceso_id)
        eq_por_persona: dict = {}
        for eq in eq_accesos_obj:
            if not eq.registro_acceso_id and eq.timestamp:
                eq_por_persona.setdefault(eq.persona_id, []).append(eq)

        # ── Serializar personas ────────────────────────────────────────
        personas_list = [{
            "id": p.id, "nombre": p.nombre, "tipo_doc": p.tipo_doc,
            "numero_doc": p.numero_doc, "email": p.email or "",
            "telefono": p.telefono or "", "perfil": p.perfil,
            "programa": p.programa or "", "ficha": p.ficha or "",
            "especialidad": p.especialidad or "", "area": p.area or "",
            "verificado": p.verificado,
            "created_at": p.created_at.isoformat() if p.created_at else "",
        } for p in personas_obj]

        computadores_list = [{
            "id": c.id, "codigo_barras": c.codigo_barras,
            "serie": c.serie or "", "marca": c.marca, "modelo": c.modelo,
            "tipo": c.tipo, "estado": c.estado,
            "asignado_a_id": c.asignado_a_id or "",
            "ubicacion": c.ubicacion or "",
            "created_at": c.created_at.isoformat() if c.created_at else "",
        } for c in computadores_obj]

        # ── Resolver equipos por acceso (directo + fallback ±30 s) ────
        def resolver_equipos(rid, pid, ts):
            eqs = eq_por_acceso.get(rid, [])
            if not eqs and pid and ts:
                eqs = [
                    e for e in eq_por_persona.get(pid, [])
                    if e.timestamp and abs((e.timestamp - ts).total_seconds()) <= 30
                ]
            if not eqs:
                return "", ""
            codigos, nombres = [], []
            for e in eqs:
                obj = equipos_map.get(e.computador_id)
                if obj:
                    codigos.append(obj.codigo_barras)
                    nombres.append(f"{obj.marca} {obj.modelo}")
            return " / ".join(codigos), " / ".join(nombres)

        # ── Construir filas enriquecidas ───────────────────────────────
        accesos_enriquecidos = []
        for r in accesos_obj:
            persona = personas_doc_map.get(r.numero_doc) or personas_map.get(r.persona_id)
            fecha_str = hora_str = ""
            if r.timestamp:
                fecha_str = r.timestamp.strftime("%d/%m/%Y")
                hora_str  = r.timestamp.strftime("%H:%M:%S")
            eq_cod, eq_nom = resolver_equipos(r.id, r.persona_id, r.timestamp)
            accesos_enriquecidos.append({
                "fecha":      fecha_str,
                "hora":       hora_str,
                "nombre":     persona.nombre   if persona else r.numero_doc,
                "numero_doc": r.numero_doc,
                "tipo_doc":   persona.tipo_doc if persona else "",
                "perfil":     persona.perfil   if persona else "",
                "tipo":       r.tipo,
                "ambiente":   r.ambiente or "",
                "eq_codigo":  eq_cod,
                "eq_nombre":  eq_nom,
                "programa":   persona.programa if persona else "",
                "ficha":      persona.ficha    if persona else "",
            })

        # ── Estadísticas ───────────────────────────────────────────────
        hoy_iso = datetime.utcnow().date().isoformat()
        hoy_fmt = datetime.now().strftime("%d/%m/%Y")
        entradas_hoy   = sum(1 for r in accesos_obj
                             if r.tipo == "ENTRADA" and r.timestamp
                             and r.timestamp.date().isoformat() == hoy_iso)
        salidas_hoy    = sum(1 for r in accesos_obj
                             if r.tipo == "SALIDA" and r.timestamp
                             and r.timestamp.date().isoformat() == hoy_iso)
        con_equipo_hoy = sum(1 for a in accesos_enriquecidos
                             if a["fecha"] == hoy_fmt and a["eq_codigo"])
        perfiles = [p.perfil.upper() for p in personas_obj]
        stats = {
            "total_personas":     len(personas_list),
            "total_accesos":      len(accesos_enriquecidos),
            "total_computadores": len(computadores_list),
            "total_eq_accesos":   len(eq_accesos_obj),
            "aprendices":   perfiles.count("APRENDIZ"),
            "instructores": perfiles.count("INSTRUCTOR"),
            "visitantes":   perfiles.count("VISITANTE"),
            "otros":        sum(1 for x in perfiles
                                if x not in ("APRENDIZ","INSTRUCTOR","VISITANTE")),
            "entradas_hoy":   entradas_hoy,
            "salidas_hoy":    salidas_hoy,
            "con_equipo_hoy": con_equipo_hoy,
        }

        # ── Construir libro ────────────────────────────────────────────
        fecha_gen = datetime.now().strftime("%d/%m/%Y  %H:%M:%S")
        wb = Workbook()
        wb.remove(wb.active)

        hoja_diario(wb, accesos_enriquecidos, fecha_gen)
        hoja_resumen_dia(wb, accesos_enriquecidos, stats, fecha_gen)
        hoja_personas(wb, personas_list, fecha_gen)
        hoja_equipos(wb, computadores_list, personas_por_id, fecha_gen)

        wb.properties.title   = "Control de Acceso SENA"
        wb.properties.subject = "Reporte oficial de ingresos y salidas"
        wb.properties.creator = "Sistema SENA"

        # ── Guardar ────────────────────────────────────────────────────
        buf = io.BytesIO()
        wb.save(buf)
        excel_bytes = buf.getvalue()

        ts             = datetime.now().strftime("%Y%m%d_%H%M%S")
        nombre_archivo = f"SENA_Control_Acceso_{ts}.xlsx"

        base_dir     = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        respaldo_dir = os.path.join(base_dir, "RESPALDOS_USB", f"respaldo_{ts}")
        os.makedirs(respaldo_dir, exist_ok=True)
        ruta_local   = os.path.join(respaldo_dir, nombre_archivo)
        with open(ruta_local, "wb") as f:
            f.write(excel_bytes)

        usb_guardado = False
        ruta_final   = ruta_local
        usb = detectar_usb()
        if usb:
            try:
                usb_dir    = os.path.join(usb, "SENA_Respaldos")
                os.makedirs(usb_dir, exist_ok=True)
                ruta_final = os.path.join(usb_dir, nombre_archivo)
                with open(ruta_final, "wb") as f:
                    f.write(excel_bytes)
                usb_guardado = True
            except Exception as e:
                print(f"[EXCEL] No se pudo guardar en USB: {e}")
                ruta_final = ruta_local

        return excel_bytes, nombre_archivo, usb_guardado, ruta_final
