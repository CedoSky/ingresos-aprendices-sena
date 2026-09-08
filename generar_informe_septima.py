#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de Informe del Sistema de Control de Ingresos SENA - SÉPTIMA EDICIÓN
Tablas bien formateadas con texto limpio y ordenado
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime

def add_heading_styled(doc, text, level):
    """Agrega un encabezado con estilo personalizado"""
    heading = doc.add_heading(text, level=level)
    heading.style = f'Heading {level}'
    
    if level == 1:
        for run in heading.runs:
            run.font.color.rgb = RGBColor(26, 82, 0)
            run.font.size = Pt(26)
            run.font.bold = True
    elif level == 2:
        for run in heading.runs:
            run.font.color.rgb = RGBColor(57, 169, 0)
            run.font.size = Pt(16)
    
    return heading

def add_paragraph_justified(doc, text, size=11, color=None):
    """Agrega párrafo justificado"""
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    for run in p.runs:
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    
    return p

def shade_cell(cell, color):
    """Agrega sombreado a una celda"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._element.get_or_add_tcPr().append(shading_elm)

def set_cell_vertical_alignment(cell, alignment):
    """Alinea verticalmente contenido de celda"""
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    tcVAlign = OxmlElement('w:vAlign')
    tcVAlign.set(qn('w:val'), alignment)
    tcPr.append(tcVAlign)

def create_formatted_table(doc, data, header_color='1a5200', col_widths=None):
    """Crea tabla con mejor formato de texto"""
    rows = len(data)
    cols = len(data[0]) if data else 1
    
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Light Grid Accent 1'
    table.autofit = False
    table.allow_autofit = False
    
    # Encabezado
    for col_idx in range(cols):
        header_cell = table.rows[0].cells[col_idx]
        shade_cell(header_cell, header_color)
        set_cell_vertical_alignment(header_cell, 'center')
        
        # Limpiar y reescribir encabezado
        header_cell.text = ''
        p = header_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        
        run = p.add_run(str(data[0][col_idx]))
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(11)
    
    # Datos
    for row_idx in range(1, rows):
        for col_idx in range(cols):
            cell = table.rows[row_idx].cells[col_idx]
            set_cell_vertical_alignment(cell, 'top')
            
            # Limpiar celda
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.2
            
            # Agregar contenido
            content = str(data[row_idx][col_idx])
            run = p.add_run(content)
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)
    
    # Anchos de columnas
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)
    
    return table

def main():
    doc = Document()
    
    # ==== PORTADA ====
    title = doc.add_heading('Sistema de Control de Ingresos SENA', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(26, 82, 0)
        run.font.size = Pt(34)
    
    subtitle = doc.add_paragraph('Informe Completo - Séptima Edición')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(57, 169, 0)
        run.font.bold = True
    
    doc.add_paragraph()
    
    info_text = doc.add_paragraph(
        f'Centro de Formación SENA\n'
        f'Versión: 7.0 - Tablas Optimizadas\n'
        f'Fecha: {datetime.now().strftime("%d de %B de %Y")}'
    )
    info_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in info_text.runs:
        run.font.size = Pt(11)
    
    doc.add_page_break()
    
    # ===============================================
    # MÓDULO 1: INGRESO
    # ===============================================
    
    add_heading_styled(doc, 'MÓDULO 1: SISTEMA DE INGRESO CON CÓDIGO DE BARRAS', 1)
    
    add_paragraph_justified(doc,
        'El módulo de ingreso es donde los aprendices registran su entrada y salida diaria. '
        'Sistema simple: escanear código de barras y listo.')
    
    doc.add_paragraph()
    
    # Tabla: Flujo General
    add_heading_styled(doc, 'Flujo General del Sistema de Ingreso', 2)
    
    flujo_data = [
        ['PASO', 'ACCIÓN', 'DESCRIPCIÓN', 'DURACIÓN'],
        ['1', 'Acercarse a pantalla', 'Persona va a la entrada donde está el lector de código de barras', '5 seg'],
        ['2', 'Escanear documento', 'Usa el lector para capturar el código de barras de su cédula', '2 seg'],
        ['3', 'Buscar en BD', 'Sistema consulta la base de datos para verificar si existe', '1 seg'],
        ['4', 'Decidir entrada/salida', 'Sistema determina si es ENTRADA (primer acceso) o SALIDA (se va)', 'Automático'],
        ['5', 'Guardar registro', 'Registra en BD con hora exacta, fecha, documento y foto', 'Automático'],
        ['6', 'Mostrar resultado', 'Pantalla muestra confirmación con nombre, programa y hora', '2 seg'],
    ]
    
    create_formatted_table(doc, flujo_data, '1a5200', [0.9, 1.8, 3.2, 1.1])
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Tabla: Pasos para Ingresar
    add_heading_styled(doc, 'Pasos Detallados para Ingresar al Sistema', 2)
    
    pasos_data = [
        ['PASO', 'ACCIÓN DEL USUARIO', 'QUÉ VE EN PANTALLA'],
        ['1', 'Llega a la entrada y se acerca a la computadora/tablet', 'Logo del SENA, cuadro de texto, instrucciones: "Escanea tu documento"'],
        ['2', 'Saca su cédula o carnet del bolsillo, limpia el código de barras', 'Pantalla lista esperando escaneo'],
        ['3', 'Acerca la cédula al lector de código de barras. Escucha un beep.', 'Sistema procesando, ícono de carga aparece'],
        ['4', 'Espera 1-2 segundos mientras sistema busca su documento', 'Indicador de procesamiento: "Buscando en base de datos..."'],
        ['5', 'Ve el resultado en la pantalla con su nombre y foto', 'Mensaje: ✅ INGRESO EXITOSO o ✅ SALIDA EXITOSA con hora exacta'],
        ['6', 'Procede a entrar o salir normalmente', 'Registro guardado automáticamente en BD'],
    ]
    
    create_formatted_table(doc, pasos_data, '1a5200', [0.8, 2.8, 3.4])
    doc.add_paragraph()
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Mensajes
    add_heading_styled(doc, 'Mensajes del Sistema: Qué Significan', 2)
    
    mensajes_data = [
        ['MENSAJE', 'SIGNIFICADO', 'QUÉ HACER'],
        ['✅ INGRESO EXITOSO', 'Tu documento fue encontrado y es tu primer acceso del día. Tu entrada está registrada.', 'Puedes pasar. Ya estás dentro del sistema.'],
        ['✅ SALIDA EXITOSA', 'Tu documento fue encontrado y ya habías ingresado hoy. Tu salida está registrada completando el día.', 'Completaste tu registro. Puedes retirarte.'],
        ['⚠️ NO ENCONTRADO', 'Tu documento no existe en la base de datos. Tu información aún no fue registrada en el sistema.', 'Ve a vigilancia. Ellos te registrarán manualmente.'],
        ['⚠️ ACCESO DENEGADO', 'Tu cuenta está desactivada por algún motivo: vacaciones, baja, suspensión o sanción.', 'Contacta administración para verificar tu estado.'],
        ['⚠️ ERROR DE LECTURA', 'El código de barras no se pudo leer. Está dañado, sucio, rayado o roto.', 'Limpia tu documento e intenta de nuevo. Si no funciona, ve a vigilancia.'],
        ['⚠️ ERROR DE CONEXIÓN', 'Problema técnico: sin internet o servidor caído. Sistema no puede consultar la BD.', 'Intenta de nuevo en 10 segundos. Si sigue fallando, reporta a vigilancia.'],
    ]
    
    create_formatted_table(doc, mensajes_data, '1a5200', [1.3, 2.3, 2.9])
    doc.add_paragraph()
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Requisitos
    add_heading_styled(doc, 'Requisitos Necesarios para que Funcione', 2)
    
    requisitos_data = [
        ['REQUISITO', 'QUÉ SIGNIFICA', 'SO FALLA, QUÉ PASA'],
        ['Estar registrado en BD', 'Tu información (nombre, documento, programa) debe estar cargada en base de datos', 'No puedes ingresar, te dice "no encontrado"'],
        ['Código de barras legible', 'Tu cédula debe estar en buen estado, sin daños, manchas ni roturas', 'El lector no puede capturarlo, necesitas registro manual'],
        ['Internet conectado', 'La computadora de ingreso debe tener conexión activa (WiFi o ethernet)', 'No se puede registrar acceso, queda sin conexión'],
        ['Pantalla y lector funcionales', 'Computadora, monitor, lector deben estar encendidos y trabajando correctamente', 'No puedes ingresar, están en mantenimiento'],
        ['Hora correcta', 'La computadora debe tener la hora real sincronizada automáticamente', 'Los registros aparecen con hora incorrecta en reportes'],
    ]
    
    create_formatted_table(doc, requisitos_data, '1a5200', [1.5, 2.2, 2.8])
    doc.add_paragraph()
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Problemas
    add_heading_styled(doc, 'Problemas Más Comunes y Soluciones', 2)
    
    problemas_data = [
        ['PROBLEMA', 'CAUSA POSIBLE', 'SOLUCIÓN'],
        ['El lector no escanea', 'Código sucio, rayado, dañado o lector desconectado', '1. Limpia tu cédula con paño suave\n2. Intenta de nuevo\n3. Si no funciona, ve a vigilancia para registro manual'],
        ['Documento no encontrado', 'Tu información no fue registrada aún en el sistema', '1. Acude a administración\n2. Lleva tu cédula y datos personales\n3. Te registran en el sistema'],
        ['Acceso Denegado sin razón', 'Tu cuenta fue desactivada o hay un error de estado', '1. Contacta administración\n2. Solicita revisión de tu estado\n3. Pide que reactiven tu acceso'],
        ['Pantalla no responde', 'Problema técnico, falta de conexión o sobrecarga del sistema', '1. Espera 10 segundos\n2. Intenta de nuevo\n3. Si persiste, reporta a vigilancia'],
        ['Lector sin beep/sonido', 'Dispositivo desconectado, batería agotada o falla física', '1. Reporta a vigilancia\n2. Ellos verifican el dispositivo\n3. Mientras, cuentan acceso manual'],
    ]
    
    create_formatted_table(doc, problemas_data, '1a5200', [1.4, 1.8, 3.3])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # ===============================================
    # MÓDULO 2: VIGILANCIA
    # ===============================================
    
    add_heading_styled(doc, 'MÓDULO 2: PANEL DE VIGILANCIA', 1)
    
    add_paragraph_justified(doc,
        'El Panel de Vigilancia es la pantalla de control del vigilante. Monitorea quién está adentro, '
        'detecta problemas y controla accesos en tiempo real.')
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Tabla: Cómo acceder
    add_heading_styled(doc, 'Cómo Acceder al Panel de Vigilancia', 2)
    
    acceso_data = [
        ['PASO', 'ACCIÓN', 'RESULTADO'],
        ['1', 'Abrir navegador web (Chrome, Firefox, Edge, etc.)', 'Navegador listo'],
        ['2', 'Ingresar email (vigilante01@sena.edu.co) y contraseña', 'Dashboard se abre con menú de opciones'],
        ['3', 'Buscar y hacer clic en "CONTROL DE ACCESOS" o "VIGILANCIA"', 'Se abre el Panel de Vigilancia'],
        ['4', 'Esperar 2-5 segundos a que carguen todos los datos', 'Panel completamente activo con información en tiempo real'],
        ['5', 'Ver la pantalla de monitoreo con registros, alertas y estadísticas', '✅ Sistema listo para vigilancia'],
    ]
    
    create_formatted_table(doc, acceso_data, '1a5200', [0.8, 2.5, 2.2])
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Tabla: Componentes
    add_heading_styled(doc, 'Partes del Panel de Vigilancia', 2)
    
    componentes_data = [
        ['PARTE', 'UBICACIÓN', 'MUESTRA', 'PARA QUÉ'],
        ['BARRA SUPERIOR (Topbar)', 'Arriba de la pantalla', 'Logo SENA, nombre panel, RELOJ DIGITAL con hora, nombre vigilante, botón SALIR', 'Verificar hora sincronizada, identificar quién está conectado'],
        ['TARJETAS DE ESTADÍSTICAS', 'Debajo de la barra', 'Ingresos Hoy, Personas Dentro Ahora, Faltas por Salir, Alertas Activas', 'Ver resumen rápido del día, identificar anomalías'],
        ['PANEL DE ZONAS/PISOS', 'Columna izquierda', 'Distribución: Piso 1 (15 personas), Piso 2 (8), Piso 3 (5)', 'Localizar dónde están personas, detectar aglomeraciones'],
        ['TABLA PRINCIPAL', 'Centro de pantalla (mayoría del espacio)', 'Listado de accesos: hora, nombre, documento, tipo (entrada/salida), estado, foto', 'Ver cada movimiento, buscar personas específicas'],
        ['FILTROS Y BÚSQUEDA', 'Arriba de la tabla', 'Botones: ENTRADA/SALIDA/TODOS, campos de búsqueda por nombre o documento', 'Filtrar información, encontrar datos rápidamente'],
    ]
    
    create_formatted_table(doc, componentes_data, '1a5200', [1.2, 1.5, 2.5, 2.3])
    doc.add_paragraph()
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Monitor tiempo real
    add_heading_styled(doc, 'Cómo Funciona el Monitor en Tiempo Real', 2)
    
    realtime_data = [
        ['CARACTERÍSTICA', 'QUÉ OCURRE', 'POR QUÉ IMPORTA'],
        ['Cada nueva entrada aparece en la tabla', 'Fila nueva se ilumina con color verde por 3 segundos indicando registro nuevo', 'El vigilante ve inmediatamente quién llegó, sin necesidad de buscar'],
        ['Color verde para entradas', 'Las personas que ENTRAN aparecen con fondo verde o ícono de flecha hacia adentro', 'Distinguir rápidamente quién está llegando de quién se va'],
        ['Color naranja para salidas', 'Las personas que SALEN aparecen con fondo naranja o ícono de flecha hacia afuera', 'Identificar quién se retira, controlar salidas al final del día'],
        ['Registros más recientes primero', 'La tabla siempre ordenada: lo más nuevo en la parte superior visible', 'Los eventos nuevos siempre están a la vista, sin necesidad de scroll'],
        ['Se actualiza automáticamente', 'Cada 2 segundos, la pantalla se refresca con nuevos datos', 'Información siempre actualizada, sin necesidad de hacer refresh (F5)'],
        ['Clic para ver detalles', 'Al hacer clic en una fila, aparece modal con foto, programa, email, historial', 'Obtener información detallada de una persona rápidamente'],
    ]
    
    create_formatted_table(doc, realtime_data, '1a5200', [1.6, 2.2, 2.7])
    doc.add_paragraph()
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Alertas
    add_heading_styled(doc, 'Sistema de Alertas: Detectar Problemas', 2)
    
    alertas_data = [
        ['TIPO DE ALERTA', 'QUÉ SIGNIFICA', 'ACCIÓN DEL VIGILANTE'],
        ['🚨 Acceso Denegado', 'Alguien intentó entrar pero fue RECHAZADO por cuenta suspendida o no autorizada', 'Interrogar a la persona, verificar identidad, avisar a administración'],
        ['⚠️ Documento No Encontrado', 'Persona escaneó pero no está registrada en el sistema, es desconocida', 'Acercarse, tomar datos, registrar como visitante temporal, avisar admin'],
        ['⏰ Acceso Fuera de Horario', 'Persona entró antes de apertura o después de cierre (fuera horarios permitidos)', 'Verificar si es personal autorizado o revisar configuración de horarios'],
        ['👥 Aglomeración Anormal', 'Un piso tiene demasiada gente, supera capacidad máxima permitida', 'Hacer ronda visual, verificar que no haya peligro, limitar nuevas entradas'],
        ['🚪 Puerta Abierta Demasiado', 'Si hay cerradura electrónica, se alerta si puerta permanece abierta anormalmente', 'Verificar seguridad de la puerta, cerrar si es necesario'],
        ['❌ Falta de Salida', 'Persona que entró hace 8+ horas aún no salió, puede estar olvidada adentro', 'Contactar por teléfono, verificar que esté bien, registrar salida'],
    ]
    
    create_formatted_table(doc, alertas_data, '1a5200', [1.5, 2.3, 2.7])
    doc.add_paragraph()
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Tareas Vigilante
    add_heading_styled(doc, 'Tareas Diarias del Vigilante', 2)
    
    tareas_data = [
        ['PERÍODO', 'HORAS', 'TAREAS PRINCIPALES', 'VERIFICAR'],
        ['🌅 MAÑANA (Apertura)', '6:00-8:30', 'Llegar temprano, abrir sistema, verificar dispositivos, recibir personas', 'Lector funciona, pantalla enciende, internet activo, hora correcta'],
        ['☀️ DURANTE DÍA (Monitoreo)', '8:30-12:00', 'Vigilancia continua, monitorear panel, revisar alertas, responder consultas', 'Accesos normales, hay alertas, personas no autorizadas, aglomeraciones'],
        ['🍽️ ALMUERZO (Zona Baja)', '12:00-1:00', 'Vigilancia reducida, revisar alertas, hacer rondas, control visual', 'Movimientos anormales, puertas seguras, falta alguien salir'],
        ['🌤️ TARDE (Salidas)', '1:00-5:00', 'Control de salidas ordenadas, vigilancia de zona, preparar reporte del día', 'Todos salen en orden, hay curiosidades, personas retenidas,evento especial'],
        ['🌆 CIERRE (Finalización)', '5:00-6:00', 'Verificar todos salieron, generar reporte, documentar eventos, cerrar sesión', 'Edificio vacío, puertas cerradas, luces apagadas, incidentes registrados'],
    ]
    
    create_formatted_table(doc, tareas_data, '1a5200', [1.1, 0.9, 2.2, 2.3])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # ===============================================
    # MÓDULO 3: ADMINISTRACIÓN
    # ===============================================
    
    add_heading_styled(doc, 'MÓDULO 3: PANEL DE ADMINISTRACIÓN', 1)
    
    add_paragraph_justified(doc,
        'El Panel de Administración es el centro de control total del sistema. Gestión de usuarios, '
        'registros de personas, reportes, configuración y auditoría completa.')
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Tabla: Tipos Admin
    add_heading_styled(doc, 'Tipos de Administradores y Sus Permisos', 2)
    
    admin_data = [
        ['TIPO', 'ACCESO TIENE', 'PUEDE HACER', 'NO PUEDE HACER'],
        ['👑 SUPERADMIN (Principal)', 'Acceso completo a TODO sin restricción. Logs, config, respaldos, auditoría.', 'Crear/editar/eliminar usuarios, cambiar configuración, eliminar datos permanentemente, restaurar respaldos', 'Nada - tiene poder total'],
        ['🛡️ ADMIN NORMAL (Regular)', 'Acceso limitado. Usuarios, reportes, algunos parámetros. Auditoría parcial.', 'Crear usuarios, editar datos, ver reportes, cambiar parámetros básicos, registrar personas', 'Eliminar usuarios, ver auditoría completa, restaurar respaldos, cambios críticos'],
    ]
    
    create_formatted_table(doc, admin_data, '1a5200', [1.3, 2.0, 2.0, 1.7])
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Tabla: Gestión Usuarios
    add_heading_styled(doc, 'Gestión de Usuarios: Crear, Editar, Contraseña', 2)
    
    gestion_data = [
        ['ACCIÓN', 'PASOS', 'CAMPOS A LLENAR', 'RESULTADO'],
        ['📝 CREAR USUARIO', '1. Click "Crear Usuario"\n2. Llenar formulario\n3. Click "GUARDAR"', 'Email único, Contraseña fuerte, Nombre, Rol (Vigilante/Admin), Activo Sí/No', 'Usuario creado, recibe email, puede iniciar sesión'],
        ['✏️ EDITAR USUARIO', '1. Buscar en lista\n2. Click "EDITAR"\n3. Modificar campos\n4. Click "GUARDAR"', 'Nombre, Email, Rol, Estado activo/inactivo (NO contraseña aquí)', 'Cambios aplicados, registrado en auditoría'],
        ['🔑 RESETEAR PSW', '1. Seleccionar usuario\n2. Click "RESETEAR CONTRASEÑA"\n3. Sistema envía email\n4. Usuario crea nueva psw', 'Sistema automáticamente envía enlace de recuperación', 'Usuario recupera acceso, contraseña anterior anulada'],
        ['🔴 DESACTIVAR CUENTA', '1. Editar usuario\n2. Cambiar "Activo" a NO\n3. Click "GUARDAR"', 'Cambiar estado a INACTIVO, guardar (tipo de desactivación temporal)', 'Usuario NO puede loguear, cuenta preservada, puede reactivarse'],
        ['🗑️ ELIMINAR (SuperAdmin)', '1. Click usuario\n2. Click "ELIMINAR"\n3. Confirmar en diálogo\n4. ⚠️ PERMANENTE', 'SuperAdmin only, requiere confirmación, no se puede deshacer', 'Usuario eliminado para siempre, registrado en auditoría'],
    ]
    
    create_formatted_table(doc, gestion_data, '1a5200', [1.2, 1.8, 1.8, 1.7])
    doc.add_paragraph()
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Tipos Personas
    add_heading_styled(doc, 'Registro de Personas: Cuatro Tipos Diferentes', 2)
    
    personas_data = [
        ['TIPO', 'QUIÉN ES', 'DATOS REQUERIDOS', 'NOTA ESPECIAL'],
        ['👨‍🎓 APRENDIZ', 'Estudiante en programa de formación', 'Nombre, Documento, Email, Teléfono, Programa, Ficha, Especialidad, Foto', 'Es 90% de registros, puede escanear código, ver historial'],
        ['👨‍🏫 INSTRUCTOR', 'Personal docente y formador', 'Nombre, Documento, Email, Teléfono, Departamento, Especialidad, Foto', 'Acceso a más áreas, horarios flexibles, reportes especiales'],
        ['👔 ADMINISTRATIVO', 'Secretaria, mantenimiento, servicios generales', 'Nombre, Documento, Email, Teléfono, Cargo, Departamento, Horario, Foto', 'Acceso a todo edificio, horarios extendidos'],
        ['👨‍💼 VISITANTE', 'Personas de afuera (proveedor, padre, etc.)', 'Nombre, Documento, Email, Teléfono, Motivo, Fecha Inicio, Fecha FIN', 'Acceso acotado con fecha límite, auto-bloqueo al vencer'],
    ]
    
    create_formatted_table(doc, personas_data, '1a5200', [1.2, 1.5, 2.0, 1.8])
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Tabla: Pasos Registro
    add_heading_styled(doc, 'Pasos para Registrar una Persona Nueva', 2)
    
    registro_data = [
        ['PASO', 'QUÉ HACER', 'DETALLES IMPORTANTES'],
        ['1', 'Acceder a REGISTRO DE PERSONAS', 'Dashboard → Click en "REGISTRO DE PERSONAS" o "+ REGISTRAR"'],
        ['2', 'Llenar datos básicos (obligatorios *)', 'Nombre completo, Tipo documento, Número documento, Email, Teléfono'],
        ['3', 'Seleccionar perfil del dropdown', 'Elegir: Aprendiz, Instructor, Administrativo o Visitante'],
        ['4', 'Llenar datos específicos según perfil', 'Aparecen campos adicionales según tipo elegido (programa, especialidad, etc.)'],
        ['5', 'Subir foto (recomendado)', 'Click "SUBIR FOTO" o arrastra. Formatos: JPG, PNG. Tamaño: 200x200px'],
        ['6', 'Click GUARDAR', 'Persona cargada en BD, ya puede escanear código, recibe email opcional'],
    ]
    
    create_formatted_table(doc, registro_data, '1a5200', [0.6, 1.8, 3.1])
    doc.add_paragraph()
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Búsqueda
    add_heading_styled(doc, 'Búsqueda, Edición y Control de Personas', 2)
    
    busqueda_data = [
        ['OPERACIÓN', 'CÓMO HACERLO', 'RESULTADO', 'IMPORTANTE'],
        ['🔍 BUSCAR', 'Campo búsqueda: nombre, documento o email. Sistema filtra instantáneamente.', 'Muestra nombre, documento, perfil, estado (activo/inactivo)', 'Búsqueda registrada en auditoría'],
        ['📋 VER DETALLES', 'Hacer clic en resultado. Se abre overlay con información completa.', 'Datos personales, foto, historial de accesos, cambios previos', 'Visualización registrada'],
        ['✏️ EDITAR', '1. Buscar persona\n2. Click "EDITAR"\n3. Cambiar campos\n4. Click "GUARDAR"', 'Campos editables: nombre, email, teléfono, programa, foto, estado', 'Cambio registrado con quién y cuándo'],
        ['🚫 DESACTIVAR', '1. Editar persona\n2. Cambiar "Activo" a NO\n3. Guardar', 'Persona no puede escanear, código rechazado, cuenta bloqueada', 'Se puede reactivar, no se elimina permanentemente'],
        ['🔑 VISITANTE (Autorizar)', 'Registrar con Perfil VISITANTE y establecer fecha inicio/fin', 'Permite acceso hasta fecha fin, después se auto-bloquea automáticamente', 'Sin intervención admin, automático'],
    ]
    
    create_formatted_table(doc, busqueda_data, '1a5200', [1.3, 2.0, 2.2, 1.5])
    doc.add_paragraph()
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Reportes
    add_heading_styled(doc, 'Generación de Reportes: Tipos Principales', 2)
    
    reportes_data = [
        ['TIPO DE REPORTE', 'INFORMACIÓN QUE CONTIENE', 'PARA QUÉ SIRVE', 'FORMATO'],
        ['📅 ASISTENCIA DIARIA', 'Quién entró/salió, horas, total trabajado, estado presencia', 'Control de asistencia, cálculo horas, justificar faltas', 'Excel, PDF'],
        ['📊 RESUMEN MENSUAL', 'Mes completo, asistencia, faltas, retrasos, salidas temprano', 'Análisis mensual, recursos, ocupación promedio', 'Excel, PDF'],
        ['👥 PERSONAS REGISTRADAS', 'Lista completa con nombre, documento, perfil, email, teléfono', 'Control de población, auditoría general, comunicaciones', 'Excel'],
        ['📍 POR ZONA/PISO', 'Accesos por zona, frecuencia, horarios, patrones de uso', 'Seguridad por ubicación, ocupación, optimizar espacios', 'Excel'],
        ['⚠️ ALERTAS Y ANOMALÍAS', 'Accesos denegados, documentos no encontrados, fallos seguridad', 'Investigación, seguridad, reporte a directivas', 'Excel, PDF'],
        ['💰 INGRESOS (si aplica)', 'Pagos por servicios (cafetería, fotocopia), totales por día/mes', 'Finanzas, auditoría contable, análisis económico', 'Excel'],
    ]
    
    create_formatted_table(doc, reportes_data, '1a5200', [1.4, 2.1, 1.7, 1.3])
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Tabla: Descargar Reportes
    add_heading_styled(doc, 'Cómo Generar y Descargar un Reporte', 2)
    
    download_data = [
        ['PASO', 'ACCIÓN A REALIZAR', 'DETALLES'],
        ['1', 'Click en "REPORTES" del menú admin', 'Opción: "REPORTES" o "GENERAR REPORTE"'],
        ['2', 'Elegir tipo de reporte', 'Dropdown: Asistencia, Mensual, Personas, Zonas, Alertas, Ingresos'],
        ['3', 'Seleccionar rango de fechas', 'Desde [fecha] Hasta [fecha]. Ejemplo: 01/04/2026 al 30/04/2026'],
        ['4', 'Aplicar filtros adicionales (opcional)', 'Si necesita: persona específica, zona, tipo evento'],
        ['5', 'Click en "GENERAR"', 'Sistema procesa (5-30 segundos según tamaño de datos)'],
        ['6', 'Esperar a que termine el procesamiento', 'Aparecerá botón "DESCARGAR" cuando esté listo'],
        ['7', 'Click "DESCARGAR EXCEL" o "DESCARGAR PDF"', 'Se descarga el archivo a tu computadora'],
        ['8', 'Abrir archivo descargado', 'Se abre en Excel o Acrobat. Puedes editar, hacer gráficos, imprimir'],
    ]
    
    create_formatted_table(doc, download_data, '1a5200', [0.7, 2.0, 3.8])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Configuración
    add_heading_styled(doc, 'Configuración del Sistema', 2)
    
    config_data = [
        ['PARÁMETRO', 'QUÉ CONTROLA', 'CÓMO CAMBIAR', 'EFECTO'],
        ['🌐 IP ESP32', 'Dirección del dispositivo cerradura electrónica', 'Campo texto: "192.168.1.45", botón "Probar", click "Guardar"', 'Sistema conecta a puerta automáticamente'],
        ['🏢 PISOS/ZONAS', 'Cuántos pisos tiene el edificio y nombres', 'Click "Agregar piso", nombre, capacidad máxima, guardar', 'Aparecen en panel vigilancia automáticamente'],
        ['🕐 HORARIOS PERMITIDOS', 'A qué horas permite entrada/salida', 'Apertura: 6:00 AM, Cierre: 6:00 PM, Fuera horario: Sí/No', 'Alertas si acceso fuera de horario permitido'],
        ['📢 MENSAJES PERSONALIZADOS', 'Textos que ven aprendices en pantalla', 'Campo texto editable: "Bienvenido al SENA", "Error: Intenta..."', 'Experiencia personalizada para usuarios'],
        ['🔐 SEGURIDAD (SuperAdmin)', 'Encriptación, retención logs, respaldos automáticos', 'Opciones avanzadas: Encriptación ON/OFF, Retención 30 días', 'Protección y respaldo de datos'],
    ]
    
    create_formatted_table(doc, config_data, '1a5200', [1.2, 1.6, 2.2, 1.5])
    doc.add_paragraph()
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # Tabla: Tareas Críticas
    add_heading_styled(doc, 'Tareas Críticas Que Debe Hacer el Administrador', 2)
    
    tareas_admin_data = [
        ['FRECUENCIA', 'TAREAS A REALIZAR', 'VERIFICAR QUÉ', 'RESOLVER'],
        ['🌅 DIARIO\n(Mañana)', 'Verificar sistema activo, probar lector, revisar alertas noche anterior', 'Pantalla enciende, lector funciona, internet OK, alertas pendientes', 'Si falla algo, contactar técnico'],
        ['📅 SEMANAL', 'Revisar personas nuevas, desactivar visitantes vencidos, auditoría cambios', 'Visitantes autorizados, accesos anormales, cambios no autorizados', 'Investigar anomalía si existe'],
        ['📊 MENSUAL', 'Generar reportes, analizar estadísticas, hacer respaldo datos, reviso auditoría mes', 'Ocupación normal, alertas con patrones, respaldo completo, cambios registrados', 'Revisar y documentar resultados'],
        ['🔐 TRIMESTRAL', 'Cambiar contraseña SuperAdmin, revisar permisos usuarios, verificar respaldos, reportar', 'Usuarios con acceso correcto, respaldos funcionan, estadísticas OK', 'Resolver antes del siguiente trimestre'],
        ['🚨 EVENTOS', 'Acceso denegado: investigar, fallo: reiniciar, incidente: alertar, visitante: verificar', 'Quién es persona, hay log del evento, dirección enterada, cambio PSW', 'Documentar y seguimiento continuo'],
    ]
    
    create_formatted_table(doc, tareas_admin_data, '1a5200', [1.1, 2.0, 2.0, 1.4])
    doc.add_paragraph()
    
    doc.add_page_break()
    
    # ==== CONCLUSIÓN ====
    add_heading_styled(doc, 'Conclusión: Los Tres Módulos Trabajando Juntos', 1)
    
    resumen_data = [
        ['MÓDULO', 'QUIÉN LO USA', 'FUNCIÓN PRINCIPAL', 'HERRAMIENTA CLAVE'],
        ['🎯 INGRESO', 'Aprendices, Instructores, Visitantes', 'Escanear código y registrar entrada/salida automáticamente', 'Lector código de barras + Pantalla'],
        ['👮 VIGILANCIA', 'Vigilantes, Personal Seguridad', 'Monitorear en tiempo real, detectar alertas, controlar acceso', 'Panel tiempo real + Filtros + Alertas'],
        ['🛡️ ADMINISTRACIÓN', 'Administradores, SuperAdmin', 'Registrar personas, gestionar usuarios, generar reportes, configurar', 'Base de datos + Reportes + Auditoría'],
    ]
    
    create_formatted_table(doc, resumen_data, '1a5200', [1.3, 1.8, 2.3, 1.8])
    doc.add_paragraph()
    doc.add_paragraph()
    
    conclusion_text = doc.add_paragraph(
        'Estos tres módulos funcionan juntos en armonía: El aprendiz ingresa (módulo 1), '
        'el vigilante lo monitorea (módulo 2), y el administrador gestiona todo (módulo 3). '
        'Es un SISTEMA INTEGRADO Y COMPLETO de seguridad y control.'
    )
    conclusion_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in conclusion_text.runs:
        run.font.size = Pt(11)
        run.italic = True
        run.font.color.rgb = RGBColor(26, 82, 0)
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    footer = doc.add_paragraph('Sistema de Control de Ingresos SENA - Séptima Edición © 2026')
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(120, 120, 120)
    
    # ==== GUARDAR ====
    output_path = r'C:\Users\AMAT\Desktop\ingreso de aprendices\INFORME_SISTEMA_SEPTIMA_EDICION.docx'
    doc.save(output_path)
    print(f'✅ SÉPTIMA EDICIÓN creada exitosamente')
    print(f'📄 Archivo: INFORME_SISTEMA_SEPTIMA_EDICION.docx')
    print(f'📊 Tablas optimizadas con texto bien formateado')
    print(f'📁 Ubicación: C:\\Users\\AMAT\\Desktop\\ingreso de aprendices\\')

if __name__ == '__main__':
    main()
